import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import joinedload

from app.main import app
from app.services.auth_service import create_access_token
from app.database.postgres import SessionLocal
from app.models.sql_models import User, Course, StudyMaterial, Enrollment
from app.services.file_storage_service import FileStorageService

client = TestClient(app)


def get_token_for_user(email: str) -> str:
    db = SessionLocal()
    user = db.query(User).options(joinedload(User.role)).filter(User.email == email).first()
    assert user is not None, f"User {email} not found"
    role_name = user.role.name if user.role else "STUDENT"
    token = create_access_token({"sub": str(user.id), "email": user.email, "role": role_name})
    db.close()
    return token


@pytest.fixture(scope="module")
def teacher_auth():
    token = get_token_for_user("prof.alan@univ.edu")
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def student_auth():
    token = get_token_for_user("alex@student.univ.edu")
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def admin_auth():
    token = get_token_for_user("admin@univ.edu")
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def course_data():
    """Ensure course has enrolled student Alex and instructor Prof Alan."""
    db = SessionLocal()
    teacher = db.query(User).filter(User.email == "prof.alan@univ.edu").first()
    student = db.query(User).filter(User.email == "alex@student.univ.edu").first()
    assert teacher and student

    course = db.query(Course).filter(Course.teacher_id == teacher.id).first()
    if not course:
        course = db.query(Course).first()

    # Ensure student is enrolled in this course
    enrolled = db.query(Enrollment).filter(
        Enrollment.student_id == student.id,
        Enrollment.course_id == course.id
    ).first()
    if not enrolled:
        enrolled = Enrollment(student_id=student.id, course_id=course.id)
        db.add(enrolled)
        db.commit()

    c_id = course.id
    db.close()
    return {"course_id": c_id, "teacher_id": teacher.id, "student_id": student.id}


# --- 1. Missing Token Rejection ---
def test_download_missing_token_rejected(course_data):
    """Unauthenticated request must be rejected with 401 Unauthorized."""
    res = client.get(f"/api/courses/{course_data['course_id']}/materials/999/download")
    assert res.status_code == 401
    assert "Missing authentication credentials" in res.json()["detail"]


# --- 2. Expired / Invalid Token Rejection ---
def test_download_invalid_token_rejected(course_data):
    """Malformed or invalid JWT token must be rejected with 401 Unauthorized."""
    headers = {"Authorization": "Bearer invalid_or_expired_jwt_token_payload"}
    res = client.get(f"/api/courses/{course_data['course_id']}/materials/999/download", headers=headers)
    assert res.status_code == 401
    assert "Invalid authentication token" in res.json()["detail"]


# --- 3. Unauthorized User Access (Student not enrolled) ---
def test_download_unauthorized_student(course_data):
    """Student attempting to download material from an un-enrolled course must receive 403 Forbidden."""
    db = SessionLocal()
    other_student = db.query(User).join(User.role).filter(User.role.has(name="STUDENT"), User.email != "alex@student.univ.edu").first()
    if not other_student:
        # Create a temporary student
        from app.models.sql_models import Role
        role = db.query(Role).filter(Role.name == "STUDENT").first()
        other_student = User(email="temp_student@univ.edu", password_hash="dummy", full_name="Temp Student", role_id=role.id)
        db.add(other_student)
        db.commit()
        db.refresh(other_student)

    # Make sure this student is NOT enrolled in course_data['course_id']
    un_enroll = db.query(Enrollment).filter(
        Enrollment.student_id == other_student.id,
        Enrollment.course_id == course_data["course_id"]
    ).first()
    if un_enroll:
        db.delete(un_enroll)
        db.commit()

    other_token = create_access_token({"sub": str(other_student.id), "email": other_student.email, "role": "STUDENT"})
    db.close()

    # Create dummy material
    db = SessionLocal()
    mat = StudyMaterial(course_id=course_data["course_id"], title="Restricted Notes", content="Secret lecture content")
    db.add(mat)
    db.commit()
    mat_id = mat.id
    db.close()

    try:
        headers = {"Authorization": f"Bearer {other_token}"}
        res = client.get(f"/api/courses/{course_data['course_id']}/materials/{mat_id}/download", headers=headers)
        assert res.status_code == 403
        assert "enrolled" in res.json()["detail"].lower()
    finally:
        db = SessionLocal()
        m = db.query(StudyMaterial).filter(StudyMaterial.id == mat_id).first()
        if m:
            db.delete(m)
            db.commit()
        db.close()


# --- 4. Missing Course Handling ---
def test_download_missing_course(student_auth):
    """Accessing non-existent course ID must return 404 Not Found."""
    res = client.get("/api/courses/999999/materials/1/download", headers=student_auth)
    assert res.status_code == 404
    assert "Course not found" in res.json()["detail"]


# --- 5. Missing Material Handling ---
def test_download_missing_material(course_data, student_auth):
    """Accessing non-existent material ID must return 404 Not Found."""
    res = client.get(f"/api/courses/{course_data['course_id']}/materials/999999/download", headers=student_auth)
    assert res.status_code == 404
    assert "Study material not found" in res.json()["detail"]


# --- 6. Authenticated PDF Download & View ---
def test_download_authenticated_pdf(course_data, teacher_auth, student_auth):
    """Upload and download real PDF: returns application/pdf and proper Content-Disposition."""
    c_id = course_data["course_id"]
    pdf_bytes = b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R>>endobj\n4 0 obj<</Length 44>>stream\nBT /F1 12 Tf 72 712 Td (Database Systems PDF) Tj ET\nendstream\nendobj\nxref\n0 5\n0000000000 65535 f\n0000000009 00000 n\n0000000056 00000 n\n0000000111 00000 n\n0000000212 00000 n\ntrailer<</Size 5/Root 1 0 R>>\nstartxref\n306\n%%EOF"
    
    files = {"files": ("lecture_schema_design.pdf", pdf_bytes, "application/pdf")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        # Enrolled student downloads
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "application/pdf" in res_dl.headers["content-type"]
        assert "lecture_schema_design.pdf" in res_dl.headers.get("content-disposition", "")
        assert res_dl.content.startswith(b"%PDF-")
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 7. Authenticated DOCX Download ---
def test_download_authenticated_docx(course_data, teacher_auth, student_auth):
    """Upload and download real DOCX: returns proper Word MIME type and filename."""
    import docx
    c_id = course_data["course_id"]
    doc = docx.Document()
    doc.add_heading("ACID Transactions", level=1)
    doc.add_paragraph("Atomicity, Consistency, Isolation, and Durability ensure reliable transaction processing.")
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    files = {"files": ("acid_properties.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "officedocument.wordprocessingml.document" in res_dl.headers["content-type"]
        assert "acid_properties.docx" in res_dl.headers.get("content-disposition", "")
        assert len(res_dl.content) == len(docx_bytes)
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 8. Authenticated PPTX Download ---
def test_download_authenticated_pptx(course_data, teacher_auth, student_auth):
    """Upload and download real PPTX: returns presentation MIME type and filename."""
    import pptx
    c_id = course_data["course_id"]
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Indexing Strategies"
    bio = io.BytesIO()
    prs.save(bio)
    pptx_bytes = bio.getvalue()

    files = {"files": ("indexing_slides.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "officedocument.presentationml.presentation" in res_dl.headers["content-type"]
        assert "indexing_slides.pptx" in res_dl.headers.get("content-disposition", "")
        assert len(res_dl.content) == len(pptx_bytes)
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 9. Authenticated TXT Download & Preview ---
def test_download_authenticated_txt(course_data, teacher_auth, student_auth):
    """Upload and download TXT: returns text/plain and correct filename."""
    c_id = course_data["course_id"]
    content = "Distributed 2-Phase Commit Protocol Specification"
    files = {"files": ("two_phase_commit.txt", content.encode("utf-8"), "text/plain")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "text/plain" in res_dl.headers["content-type"]
        assert "two_phase_commit.txt" in res_dl.headers.get("content-disposition", "")
        assert content.encode("utf-8") in res_dl.content
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 10. Authenticated MD Download & Preview ---
def test_download_authenticated_md(course_data, teacher_auth, student_auth):
    """Upload and download Markdown: returns text/markdown and correct filename."""
    c_id = course_data["course_id"]
    content = "# Raft Consensus Leader Election\n\nRaft uses randomized election timers."
    files = {"files": ("raft_consensus.md", content.encode("utf-8"), "text/markdown")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "text/markdown" in res_dl.headers["content-type"]
        assert "raft_consensus.md" in res_dl.headers.get("content-disposition", "")
        assert content.encode("utf-8") in res_dl.content
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 11. Text Note Without file_storage_id (On-the-fly generated text) ---
def test_download_text_note_without_storage_id(course_data, student_auth, teacher_auth):
    """Material added as pure text (no MongoDB file_storage_id) generates dynamic text file."""
    c_id = course_data["course_id"]
    payload = {
        "title": "Query Optimizer Architecture",
        "content": "System R dynamic programming optimizer considers left-deep join trees.",
        "material_type": "NOTES"
    }
    res_add = client.post(f"/api/courses/{c_id}/materials", json=payload, headers=teacher_auth)
    assert res_add.status_code == 201
    mat_id = res_add.json()["id"]

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 200
        assert "text/plain" in res_dl.headers["content-type"]
        assert "Query Optimizer Architecture.txt" in res_dl.headers.get("content-disposition", "")
        assert b"System R dynamic programming" in res_dl.content
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 12. Corrupted / Unavailable File Storage ---
def test_download_missing_storage_file(course_data, student_auth, teacher_auth):
    """If file_storage_id points to missing document in MongoDB, must return 404 Not Found."""
    c_id = course_data["course_id"]
    db = SessionLocal()
    # Insert a material with non-existent storage ID
    mat = StudyMaterial(
        course_id=c_id,
        title="Corrupted Storage Reference",
        content="Extracted text here",
        original_filename="ghost_file.pdf",
        file_storage_id="60c72b2f9b1d8b1b8c8b4567"  # Non-existent Mongo ObjectId
    )
    db.add(mat)
    db.commit()
    mat_id = mat.id
    db.close()

    try:
        res_dl = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_dl.status_code == 404
        assert "File content not found" in res_dl.json()["detail"]
    finally:
        db = SessionLocal()
        m = db.query(StudyMaterial).filter(StudyMaterial.id == mat_id).first()
        if m:
            db.delete(m)
            db.commit()
        db.close()


# --- 13. Role-Based Access: Teacher, Admin, and Student ---
def test_download_rbac_permissions(course_data, teacher_auth, admin_auth, student_auth):
    """Instructor, Admin, and enrolled Student can all access materials."""
    c_id = course_data["course_id"]
    content = "RBAC Verification Document"
    files = {"files": ("rbac_check.txt", content.encode("utf-8"), "text/plain")}
    res_up = client.post(f"/api/courses/{c_id}/materials/upload", files=files, headers=teacher_auth)
    assert res_up.status_code == 201
    mat_id = res_up.json()[0]["id"]

    try:
        # 1. Course Teacher
        res_teacher = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=teacher_auth)
        assert res_teacher.status_code == 200

        # 2. System Administrator
        res_admin = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=admin_auth)
        assert res_admin.status_code == 200

        # 3. Enrolled Student
        res_student = client.get(f"/api/courses/{c_id}/materials/{mat_id}/download", headers=student_auth)
        assert res_student.status_code == 200
    finally:
        client.delete(f"/api/courses/{c_id}/materials/{mat_id}", headers=teacher_auth)


# --- 14. AI Search Source Material Download Link Verification ---
def test_ai_search_source_download_link_validity(course_data, student_auth):
    """AI Search returns course sources with download_url that can be fetched with auth."""
    payload = {
        "query": "What is Dijkstra shortest path algorithm?",
        "mode": "course_notes",
        "depth": "standard",
        "course_id": course_data["course_id"]
    }
    res = client.post("/api/search", json=payload, headers=student_auth)
    assert res.status_code == 200
    data = res.json()
    if data.get("course_sources") and len(data["course_sources"]) > 0:
        first_src = data["course_sources"][0]
        dl_url = first_src.get("download_url")
        assert dl_url is not None
        assert dl_url.startswith("/api/courses/")
        assert dl_url.endswith("/download")

        # Verify downloading with student auth returns real file content
        res_dl = client.get(dl_url, headers=student_auth)
        assert res_dl.status_code == 200
        assert "Content-Disposition" in res_dl.headers
        # Must not be a JSON error detail
        assert not res_dl.content.startswith(b'{"detail":')
